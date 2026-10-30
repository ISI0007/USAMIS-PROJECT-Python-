"""USAMIS four-layer deck generator. Theme matched to the existing deliverable (Arial, 16:9)."""
import json, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

THEME = {
    "navy": RGBColor(0x00, 0x2D, 0x72),
    "blue": RGBColor(0x00, 0x52, 0xCC),
    "ink": RGBColor(0x1F, 0x29, 0x33),
    "gray": RGBColor(0x55, 0x55, 0x55),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "bg": RGBColor(0xF7, 0xF9, 0xFC),
    "accent": RGBColor(0xE8, 0x6C, 0x00),
    "line": RGBColor(0xD5, 0xDD, 0xE8),
}
FONT = "Arial"
EMU_W = Inches(13.333)
EMU_H = Inches(7.5)


def _tf(shape):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    return tf


def _run(p, text, size, color, bold=False, font=FONT):
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color
    return r


def rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(1, int(x), int(y), int(w), int(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


class Deck:
    def __init__(self, theme):
        self.prs = Presentation()
        self.prs.slide_width = EMU_W
        self.prs.slide_height = EMU_H
        self.t = theme
        self.n = 0

    def save(self, path):
        self.prs.save(path)

    def _slide(self, bg=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        rect(s, 0, 0, EMU_W, EMU_H, bg or self.t["bg"])
        self.n += 1
        return s

    def _footer(self, s, label=""):
        rect(s, 0, EMU_H - Inches(0.42), EMU_W, Inches(0.42), self.t["navy"])
        tb = s.shapes.add_textbox(Inches(0.55), EMU_H - Inches(0.36), Inches(9), Inches(0.3))
        p = _tf(tb).paragraphs[0]
        _run(p, "USAMIS  ·  Jinling Institute of Technology", 10, RGBColor(0xB8, 0xC6, 0xDC))
        tb2 = s.shapes.add_textbox(Inches(11.8), EMU_H - Inches(0.36), Inches(1.2), Inches(0.3))
        p2 = _tf(tb2).paragraphs[0]; p2.alignment = PP_ALIGN.RIGHT
        _run(p2, f"{self.n:02d}", 10, RGBColor(0xB8, 0xC6, 0xDC))

    def title(self, kicker, heading, sub, meta):
        s = self._slide(self.t["navy"])
        rect(s, 0, Inches(4.6), EMU_W, Inches(0.06), self.t["accent"])
        tb = s.shapes.add_textbox(Inches(0.9), Inches(1.75), Inches(11.5), Inches(0.4))
        p = _tf(tb).paragraphs[0]
        _run(p, kicker.upper(), 13, RGBColor(0x8F, 0xB0, 0xE0), bold=True)
        tb = s.shapes.add_textbox(Inches(0.9), Inches(2.25), Inches(11.5), Inches(1.9))
        tf = _tf(tb)
        p = tf.paragraphs[0]
        _run(p, heading, 44, self.t["white"], bold=True)
        tb = s.shapes.add_textbox(Inches(0.9), Inches(4.85), Inches(11.0), Inches(0.9))
        p = _tf(tb).paragraphs[0]
        _run(p, sub, 17, RGBColor(0xC9, 0xD6, 0xEA))
        tb = s.shapes.add_textbox(Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.5))
        p = _tf(tb).paragraphs[0]
        _run(p, meta, 12, RGBColor(0x8F, 0xB0, 0xE0))
        return s

    def agenda(self, items):
        s = self._slide()
        self._head(s, "Agenda", "What this deck covers")
        y = Inches(1.95)
        for i, it in enumerate(items, 1):
            rect(s, Inches(0.85), y, Inches(0.52), Inches(0.52), self.t["blue"])
            tb = s.shapes.add_textbox(Inches(0.85), y + Inches(0.09), Inches(0.52), Inches(0.4))
            p = _tf(tb).paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            _run(p, str(i), 16, self.t["white"], bold=True)
            tb = s.shapes.add_textbox(Inches(1.62), y + Inches(0.08), Inches(10.6), Inches(0.45))
            p = _tf(tb).paragraphs[0]
            _run(p, it, 17, self.t["ink"])
            y += Inches(0.62)
        self._footer(s)
        return s

    def _head(self, s, title, sub=None):
        rect(s, Inches(0.85), Inches(0.62), Inches(0.075), Inches(0.72), self.t["accent"])
        tb = s.shapes.add_textbox(Inches(1.08), Inches(0.6), Inches(11.3), Inches(0.6))
        p = _tf(tb).paragraphs[0]
        _run(p, title, 30, self.t["navy"], bold=True)
        if sub:
            tb = s.shapes.add_textbox(Inches(1.08), Inches(1.22), Inches(11.3), Inches(0.4))
            p = _tf(tb).paragraphs[0]
            _run(p, sub, 14, self.t["gray"])
        rect(s, Inches(0.85), Inches(1.75), Inches(11.6), Inches(0.02), self.t["line"])

    def divider(self, num, title, blurb):
        s = self._slide(self.t["navy"])
        tb = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(2.2), Inches(1.6))
        p = _tf(tb).paragraphs[0]
        _run(p, num, 96, RGBColor(0x1B, 0x4A, 0x94), bold=True)
        tb = s.shapes.add_textbox(Inches(0.95), Inches(3.35), Inches(11.0), Inches(1.0))
        p = _tf(tb).paragraphs[0]
        _run(p, title, 36, self.t["white"], bold=True)
        rect(s, Inches(0.95), Inches(4.35), Inches(1.4), Inches(0.05), self.t["accent"])
        tb = s.shapes.add_textbox(Inches(0.95), Inches(4.6), Inches(10.6), Inches(1.0))
        p = _tf(tb).paragraphs[0]
        _run(p, blurb, 15, RGBColor(0xC9, 0xD6, 0xEA))
        self._footer(s)
        return s

    def bullets(self, title, sub, items, cols=1):
        s = self._slide()
        self._head(s, title, sub)
        top = Inches(2.0)
        if cols == 1:
            tb = s.shapes.add_textbox(Inches(0.95), top, Inches(11.4), Inches(4.6))
            tf = _tf(tb)
            for i, it in enumerate(items):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(11)
                _run(p, "▪  ", 15, self.t["accent"], bold=True)
                if isinstance(it, tuple):
                    _run(p, it[0] + " — ", 15, self.t["navy"], bold=True)
                    _run(p, it[1], 15, self.t["ink"])
                else:
                    _run(p, it, 15, self.t["ink"])
        else:
            half = (len(items) + 1) // 2
            for ci, chunk in enumerate([items[:half], items[half:]]):
                tb = s.shapes.add_textbox(Inches(0.95) + ci * Inches(5.85), top, Inches(5.5), Inches(4.6))
                tf = _tf(tb)
                for i, it in enumerate(chunk):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.space_after = Pt(10)
                    _run(p, "▪  ", 14, self.t["accent"], bold=True)
                    if isinstance(it, tuple):
                        _run(p, it[0] + " — ", 14, self.t["navy"], bold=True)
                        _run(p, it[1], 14, self.t["ink"])
                    else:
                        _run(p, it, 14, self.t["ink"])
        self._footer(s)
        return s

    def cards(self, title, sub, cards, cols=3):
        s = self._slide()
        self._head(s, title, sub)
        n = len(cards)
        rows = (n + cols - 1) // cols
        gap = Inches(0.32)
        cw = (Inches(11.55) - gap * (cols - 1)) / cols
        ch = Inches(1.55) if rows > 1 else Inches(2.4)
        for i, (head, body) in enumerate(cards):
            c, r = i % cols, i // cols
            x = Inches(0.9) + c * (cw + gap)
            y = Inches(1.98) + r * (ch + Inches(0.28))
            rect(s, x, y, cw, ch, self.t["white"])
            rect(s, x, y, cw, Inches(0.055), self.t["blue"])
            tb = s.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), cw - Inches(0.4), Inches(0.4))
            p = _tf(tb).paragraphs[0]
            _run(p, head, 14, self.t["navy"], bold=True)
            tb = s.shapes.add_textbox(x + Inches(0.2), y + Inches(0.66), cw - Inches(0.4), ch - Inches(0.85))
            tf = _tf(tb)
            p = tf.paragraphs[0]
            _run(p, body, 11.5, self.t["gray"])
        self._footer(s)
        return s

    def stats(self, title, sub, stats, note=None):
        s = self._slide()
        self._head(s, title, sub)
        n = len(stats)
        gap = Inches(0.3)
        cw = (Inches(11.55) - gap * (n - 1)) / n
        for i, (val, lab) in enumerate(stats):
            x = Inches(0.9) + i * (cw + gap)
            rect(s, x, Inches(2.15), cw, Inches(2.0), self.t["white"])
            tb = s.shapes.add_textbox(x, Inches(2.45), cw, Inches(0.95))
            p = _tf(tb).paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            _run(p, val, 40, self.t["blue"], bold=True)
            tb = s.shapes.add_textbox(x + Inches(0.15), Inches(3.42), cw - Inches(0.3), Inches(0.6))
            p = _tf(tb).paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            _run(p, lab, 11, self.t["gray"])
        if note:
            tb = s.shapes.add_textbox(Inches(0.9), Inches(4.5), Inches(11.5), Inches(1.2))
            tf = _tf(tb)
            p = tf.paragraphs[0]
            _run(p, note, 12, self.t["gray"])
        self._footer(s)
        return s

    def image(self, title, sub, img, caption=""):
        s = self._slide()
        self._head(s, title, sub)
        if os.path.exists(img):
            from PIL import Image
            iw, ih = Image.open(img).size
            box_w, box_h = 11.2, 4.3
            sc = min(box_w / iw, box_h / ih)
            w, h = Inches(iw * sc), Inches(ih * sc)
            x = int((EMU_W - w) / 2)
            s.shapes.add_picture(img, x, Inches(1.95), int(w), int(h))
        if caption:
            tb = s.shapes.add_textbox(Inches(0.9), Inches(6.45), Inches(11.5), Inches(0.4))
            p = _tf(tb).paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            _run(p, caption, 11, self.t["gray"])
        self._footer(s)
        return s

    def table(self, title, sub, headers, rows, widths=None):
        s = self._slide()
        self._head(s, title, sub)
        x = Inches(0.9); y = Inches(1.98)
        total = Inches(11.55)
        if widths is None:
            widths = [1.0 / len(headers)] * len(headers)
        rowh = Inches(0.42)
        rect(s, x, y, total, rowh, self.t["navy"])
        cx = x
        for i, hd in enumerate(headers):
            w = int(total * widths[i])
            tb = s.shapes.add_textbox(cx + Inches(0.12), y + Inches(0.1), w - Inches(0.2), Inches(0.3))
            p = _tf(tb).paragraphs[0]
            _run(p, hd, 11.5, self.t["white"], bold=True)
            cx += w
        ry = y + rowh
        for ri, row in enumerate(rows):
            bgc = self.t["white"] if ri % 2 == 0 else RGBColor(0xEE, 0xF2, 0xF8)
            rect(s, x, ry, total, rowh, bgc)
            cx = x
            for ci, cell in enumerate(row):
                w = int(total * widths[ci])
                tb = s.shapes.add_textbox(cx + Inches(0.12), ry + Inches(0.09), w - Inches(0.2), Inches(0.32))
                p = _tf(tb).paragraphs[0]
                _run(p, str(cell), 10.5, self.t["ink"] if ci else self.t["navy"], bold=(ci == 0))
                cx += w
            ry += rowh
        self._footer(s)
        return s

    def closing(self, heading, lines):
        s = self._slide(self.t["navy"])
        tb = s.shapes.add_textbox(Inches(0.95), Inches(2.4), Inches(11.4), Inches(1.1))
        p = _tf(tb).paragraphs[0]
        _run(p, heading, 40, self.t["white"], bold=True)
        rect(s, Inches(0.95), Inches(3.6), Inches(1.4), Inches(0.05), self.t["accent"])
        tb = s.shapes.add_textbox(Inches(0.95), Inches(3.9), Inches(11.0), Inches(2.4))
        tf = _tf(tb)
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            _run(p, ln, 15, RGBColor(0xC9, 0xD6, 0xEA))
        return s


def build(content, out):
    d = Deck(THEME)
    base = content.get("base", "")
    def res(p):
        return p if os.path.isabs(p) else os.path.join(base, p)
    for sl in content["slides"]:
        k = sl["type"]
        if k == "title":
            d.title(sl["kicker"], sl["heading"], sl["sub"], sl["meta"])
        elif k == "agenda":
            d.agenda(sl["items"])
        elif k == "divider":
            d.divider(sl["num"], sl["title"], sl["blurb"])
        elif k == "bullets":
            d.bullets(sl["title"], sl.get("sub"), sl["items"], sl.get("cols", 1))
        elif k == "cards":
            d.cards(sl["title"], sl.get("sub"), sl["cards"], sl.get("cols", 3))
        elif k == "stats":
            d.stats(sl["title"], sl.get("sub"), sl["stats"], sl.get("note"))
        elif k == "image":
            d.image(sl["title"], sl.get("sub"), res(sl["img"]), sl.get("caption", ""))
        elif k == "table":
            d.table(sl["title"], sl.get("sub"), sl["headers"], sl["rows"], sl.get("widths"))
        elif k == "closing":
            d.closing(sl["heading"], sl["lines"])
        else:
            raise ValueError(k)
    d.save(out)
    print(f"OK {len(d.prs.slides._sldIdLst)} slides -> {out}")


if __name__ == "__main__":
    build(json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2])
