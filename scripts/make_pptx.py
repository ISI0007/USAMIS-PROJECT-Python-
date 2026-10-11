"""Generate the USAMIS-Python PowerPoint presentation.

Produces docs/USAMIS_Presentation.pptx with python-pptx, embedding the exported
diagram images and diagrams. Run:

    python scripts/make_pptx.py
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DIAGRAMS = os.path.join(ROOT, "docs", "diagrams")
OUT = os.path.join(ROOT, "docs", "USAMIS_Presentation.pptx")

BRAND = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2B, 0x6C, 0xB0)
DARK = RGBColor(0x1A, 0x24, 0x33)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_title_slide(prs, title, subtitle, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _band(slide, 0, 0, prs.slide_width, Inches(4.3), BRAND)
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(12), Inches(1.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        s = tf.add_paragraph()
        s.text = subtitle
        s.font.size = Pt(18)
        s.font.color.rgb = RGBColor(0xCF, 0xDC, 0xEC)
    if footer:
        fb = slide.shapes.add_textbox(Inches(0.7), Inches(6.6), Inches(11), Inches(0.5))
        fp = fb.text_frame.paragraphs[0]
        fp.text = footer
        fp.font.size = Pt(12)
        fp.font.color.rgb = MUTED
    return slide


def _band(slide, left, top, width, height, color):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_bullets_slide(prs, title, bullets, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _band(slide, 0, 0, prs.slide_width, Inches(1.05), BRAND)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.18), Inches(12.2), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE

    top = Inches(1.4)
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.7), top, Inches(12), Inches(0.5))
        sp = sb.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)
        sp.font.italic = True
        sp.font.color.rgb = MUTED
        top = Inches(2.0)

    box = slide.shapes.add_textbox(Inches(0.7), top, Inches(12), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        para.text = ("• " if level == 0 else "– ") + text
        para.level = level
        para.font.size = Pt(18 if level == 0 else 15)
        para.font.color.rgb = DARK if level == 0 else MUTED
        para.space_after = Pt(6)
    return slide


def add_image_slide(prs, title, image_name, caption=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _band(slide, 0, 0, prs.slide_width, Inches(1.05), BRAND)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.18), Inches(12.2), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE

    path = os.path.join(DIAGRAMS, image_name)
    if os.path.exists(path):
        slide.shapes.add_picture(path, Inches(1.0), Inches(1.35), height=Inches(5.2))
    else:
        warn = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(1))
        warn.text_frame.paragraphs[0].text = f"[missing image: {image_name}]"

    if caption:
        cb = slide.shapes.add_textbox(Inches(0.7), Inches(6.7), Inches(12), Inches(0.5))
        cp = cb.text_frame.paragraphs[0]
        cp.text = caption
        cp.font.size = Pt(12)
        cp.font.color.rgb = MUTED
    return slide


def add_table_slide(prs, title, headers, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _band(slide, 0, 0, prs.slide_width, Inches(1.05), BRAND)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.18), Inches(12.2), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE

    n_rows = len(rows) + 1
    n_cols = len(headers)
    table = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5 * n_rows)
    ).table
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = header
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
        cell.text_frame.paragraphs[0].font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            cell.text_frame.paragraphs[0].font.size = Pt(13)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 — Title
    add_title_slide(
        prs,
        "University Student Academic Management\nInformation System (USAMIS)",
        "A Role-Based Management Information System using Python, Flask & PostgreSQL",
        "USAMIS-Python · Version 1.0 · [Your Name] · [Student ID] · [Date]",
    )

    # 2 — Problem statement
    add_bullets_slide(prs, "Problem Statement", [
        "Student records are scattered across paper and spreadsheets.",
        "No single authenticated entry point for staff roles.",
        "Duplicate and inconsistent data across departments.",
        "Management reports are slow to produce and error-prone.",
        "Weak control over who may view or edit records.",
        "Fee tracking and result processing are manual and unreliable.",
    ])

    # 3 — Objectives
    add_bullets_slide(prs, "Project Objectives", [
        "Centralise all student academic records in one database.",
        "Provide one login for all roles with username + password.",
        "Enforce Role-Based Access Control (RBAC) on the server.",
        "Manage students, courses, enrollment, grades, attendance, fees.",
        "Prevent duplicate records and unauthorized access.",
        "Generate live management reports from real data.",
        "Maintain an audit log of sensitive actions.",
    ])

    # 4 — Scope and roles
    add_bullets_slide(prs, "Scope & User Roles", [
        "Administrator — full control, users, audit log.",
        "Academic staff / Registrar — students, courses, enrollment.",
        "Teacher — grades and attendance for assigned courses.",
        "Finance officer — fee assessments and payments.",
        "Student — own profile, enrollment, grades, fees.",
        ("Out of scope: real payment gateway, admissions, AI/ML features.", 1),
    ])

    # 5 — Technology stack
    add_table_slide(
        prs, "Technology Stack",
        ["Layer", "Technology"],
        [
            ["Language", "Python 3.11"],
            ["Framework", "Flask (application factory, blueprints)"],
            ["ORM", "SQLAlchemy 2.0 / Flask-SQLAlchemy"],
            ["Migrations", "Flask-Migrate / Alembic"],
            ["Database", "PostgreSQL 17"],
            ["Auth / Forms", "Flask-Login, Flask-WTF, Werkzeug hashing"],
            ["UI", "HTML5, CSS3, vanilla JavaScript"],
            ["Testing", "pytest, pytest-cov"],
        ],
    )

    # 6 — System architecture
    add_image_slide(prs, "System Architecture",
                    "02-system-architecture.png",
                    "Figure 2 — layered architecture (blueprints, services, algorithms, data).")

    # 7 — ER diagram
    add_image_slide(prs, "Database ER Diagram",
                    "03-er-diagram.png",
                    "Figure 3 — entity relationship diagram of the implemented schema.")

    # 8 — Enrollment flowchart
    add_image_slide(prs, "Enrollment Process Flowchart",
                    "04-enrollment-flowchart.png",
                    "Figure 4 — Algorithm 1 with explicit rejection paths.")

    # 9 — Three algorithms
    add_bullets_slide(prs, "Three Explainable Algorithms", [
        "1. Course Enrollment Eligibility — status, duplicates, prerequisites, capacity.",
        "2. Academic Performance — credit-weighted average of published grades.",
        "3. Fee Balance — assessments − adjustments − payments (decimal-safe).",
        ("Deterministic and unit-tested — not artificial intelligence.", 1),
        ("Each returns structure: reason codes, notes, and computed detail.", 1),
    ])

    # 10 — Data flow diagram
    add_image_slide(prs, "Data Flow Diagram",
                    "05-data-flow-diagram.png",
                    "Figure 5 — movement of requests and academic data between actors, processes, and stores.")

    # 11 — System context
    add_image_slide(prs, "System Context",
                    "01-system-context.png",
                    "Figure 1 — USAMIS and its external actors and database.")

    # 12 — Security and access control
    add_bullets_slide(prs, "Security & Access Control", [
        "Passwords stored as salted hashes (Werkzeug).",
        "CSRF protection on all state-changing forms.",
        "Server-side RBAC on every protected endpoint.",
        "Ownership checks: students see only their own records.",
        "Login throttling plus account lockout after failures.",
        "Parameterized ORM queries — no SQL string concatenation.",
        "Audit logging of sensitive actions.",
        "Security headers; no debug mode in production.",
    ])

    # 13 — Testing
    add_bullets_slide(prs, "Testing Methodology & Results", [
        "pytest over a dedicated PostgreSQL test database.",
        "Unit, integration, and security test categories.",
        "Algorithms covered by dedicated unit tests.",
        "Negative security tests for access control.",
        "Actual result: 87 passed, 0 failed.",
        ("Six defects found and fixed during testing.", 1),
    ])

    # 14 — Limitations and future work
    add_bullets_slide(prs, "Limitations & Future Work", [
        "Limitations:",
        ("Fee module records payments; it is not a payment gateway.", 1),
        ("In-memory rate-limit storage by default.", 1),
        "Future work:",
        ("Mobile client; email/SMS notifications.", 1),
        ("Payment-gateway integration; QR attendance.", 1),
        ("Cloud deployment with a shared rate-limit backend.", 1),
    ])

    # 15 — Conclusion
    add_bullets_slide(prs, "Conclusion", [
        "Centralised, role-based academic management information system.",
        "Server-enforced access control with an audit trail.",
        "Three explainable algorithms, unit-tested.",
        "Tested (87 passing), documented, and reproducible.",
        "Delivered as working software with five verified diagrams.",
        ("Thank you — Questions & Demonstration.", 1),
    ])

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print(f"Saved {OUT} with {len(prs.slides.__iter__.__self__._sldIdLst)} slides")


if __name__ == "__main__":
    build()
