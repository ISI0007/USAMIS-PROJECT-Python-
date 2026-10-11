"""Verify the generated PPTX opens and its contents are readable."""
import os
import sys
import zipfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

PPTX = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "docs", "USAMIS_Presentation.pptx")
)

# 1. It must be a valid OOXML zip containing the expected parts.
assert os.path.exists(PPTX), "presentation file missing"
assert zipfile.is_zipfile(PPTX), "not a valid zip/OOXML package"
with zipfile.ZipFile(PPTX) as z:
    names = z.namelist()
    assert "[Content_Types].xml" in names, "missing content types"
    slide_parts = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    media = [n for n in names if n.startswith("ppt/media/")]
print(f"OOXML OK — {len(slide_parts)} slide parts, {len(media)} embedded media files")

# 2. Re-open with python-pptx and read the slides back.
from pptx import Presentation  # noqa: E402

prs = Presentation(PPTX)
print(f"python-pptx opened presentation with {len(prs.slides)} slides")
for i, slide in enumerate(prs.slides, start=1):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            texts.append(shape.text_frame.text.strip().splitlines()[0])
    print(f"  Slide {i:>2}: {texts[0][:70] if texts else '(no text)'}")

assert len(prs.slides) >= 12, "expected at least 12 slides"
print("\nPPTX VERIFIED")
